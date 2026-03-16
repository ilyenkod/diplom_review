"""
Процессор для обработки DOCX файлов.

Извлекает текст из документов Microsoft Word (.docx).
"""

import asyncio
from io import BytesIO
from pathlib import Path

from docx import Document

from app.infrastructure.logging import AppLogger
from app.infrastructure.processors.base import BaseProcessor, ProcessorResult


class DocxProcessor(BaseProcessor):
    """Процессор для DOCX файлов."""

    def __init__(self, logger: AppLogger) -> None:
        """Инициализирует DOCX процессор.

        Args:
            logger: Логгер приложения.
        """
        self._logger = logger

    def get_supported_extensions(self) -> set[str]:
        """Возвращает множество поддерживаемых расширений файлов.

        Returns:
            Множество поддерживаемых расширений.
        """
        return {"docx"}

    async def process(self, file_path: Path) -> ProcessorResult:
        """Обрабатывает DOCX документ и извлекает текст.

        Args:
            file_path: Путь к файлу.

        Returns:
            Результат обработки с текстом или ошибкой.
        """
        try:
            text = await asyncio.to_thread(self._extract_text_sync, file_path)
            self._logger.info(
                "DOCX file processed successfully",
                extra={"file_path": str(file_path), "text_length": len(text)},
            )
            return ProcessorResult(
                success=True,
                text=text,
                metadata={"format": "docx", "paragraphs": text.count("\n") + 1},
            )
        except Exception as e:
            self._logger.error(
                "Failed to process DOCX file",
                extra={"file_path": str(file_path), "error": str(e)},
            )
            return ProcessorResult(
                success=False,
                text="",
                error=str(e),
            )

    async def process_bytes(self, content: bytes) -> ProcessorResult:
        """Обрабатывает DOCX документ из байтов и извлекает текст.

        Args:
            content: Содержимое файла в байтах.

        Returns:
            Результат обработки с текстом или ошибкой.
        """
        try:
            text = await asyncio.to_thread(self._extract_text_from_bytes_sync, content)
            self._logger.info(
                "DOCX bytes processed successfully",
                extra={"content_length": len(content), "text_length": len(text)},
            )
            return ProcessorResult(
                success=True,
                text=text,
                metadata={"format": "docx", "paragraphs": text.count("\n") + 1},
            )
        except Exception as e:
            self._logger.error(
                "Failed to process DOCX bytes",
                extra={"content_length": len(content), "error": str(e)},
            )
            return ProcessorResult(
                success=False,
                text="",
                error=str(e),
            )

    def _extract_text_sync(self, file_path: Path) -> str:
        """Синхронное извлечение текста из файла.

        Args:
            file_path: Путь к файлу.

        Returns:
            Извлеченный текст.
        """
        doc = Document(str(file_path))
        paragraphs = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                paragraphs.append(paragraph.text)

        # Также извлекаем текст из таблиц
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    paragraphs.append(" | ".join(row_text))

        return "\n".join(paragraphs)

    def _extract_text_from_bytes_sync(self, content: bytes) -> str:
        """Синхронное извлечение текста из байтов.

        Args:
            content: Содержимое файла в байтах.

        Returns:
            Извлеченный текст.
        """
        file_obj = BytesIO(content)
        doc = Document(file_obj)
        paragraphs = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                paragraphs.append(paragraph.text)

        # Также извлекаем текст из таблиц
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    paragraphs.append(" | ".join(row_text))

        return "\n".join(paragraphs)
